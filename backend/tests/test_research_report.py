import io
import json
import re
import zipfile
from datetime import timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.main import app
from app.routers import research as research_router
from app.services import research_orchestrator as ro
from app.services import research_report as rr
from app.services import research_synthesis as rs
from test_research_orchestrator import (  # noqa: F401 - shared fixtures
    A, B, COMPLEX, COVERAGE_Q, execute, fresh_registry, make_plan, overpass,
)

client = TestClient(app)
PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def no_llm(monkeypatch):
    monkeypatch.setattr(research_router, "create_provider", lambda settings: None)


def prepare(question=COVERAGE_Q, b=None, insights=True):
    events = execute(make_plan(question, A, b))
    xid = events[0]["execution_id"]
    if insights:
        assert client.post(f"/api/research/executions/{xid}/synthesize", json={}).status_code == 200
    return xid


def report(xid, fmt="pdf"):
    return client.post(f"/api/research/executions/{xid}/report", json={"format": fmt})


def pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() for page in reader.pages)


def docx_document(content: bytes):
    return Document(io.BytesIO(content))


def docx_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    return "\n".join(parts)


def synthesis_of(xid):
    return ro.REGISTRY.get(xid).syntheses["auto"]


# Formats, MIME types and filenames

def test_pdf_is_a_valid_document_with_all_sections(overpass):
    xid = prepare(COMPLEX, b=B)
    response = report(xid, "pdf")
    assert response.status_code == 200 and response.headers["content-type"] == PDF_TYPE
    assert response.content.startswith(b"%PDF-") and b"%%EOF" in response.content[-1024:]
    reader = PdfReader(io.BytesIO(response.content))
    assert len(reader.pages) >= 5 and reader.metadata.author == "GeoMind AI"
    text = pdf_text(response.content)
    headings = ["Executive Summary", "Research Objectives", "Methodology", "Key Findings", "Comparative Analysis",
                "Site Suitability", "Evidence Register", "Data Quality and Limitations", "Conclusions and Next Steps",
                "Sources and Attribution"]
    positions = [text.index(f"{i}. {h}") for i, h in enumerate(headings, start=1)]
    assert positions == sorted(positions)
    assert f"Page 2 of {len(reader.pages)}" in text


def test_docx_is_a_valid_editable_document(overpass):
    xid = prepare(COMPLEX, b=B)
    response = report(xid, "docx")
    assert response.status_code == 200 and response.headers["content-type"] == DOCX_TYPE
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert "word/document.xml" in archive.namelist() and archive.testzip() is None
    document = docx_document(response.content)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings[0] == "1. Executive Summary" and headings[-1] == "10. Sources and Attribution"
    assert len(document.tables) >= 8
    assert "GeoMind AI" in document.sections[0].header.paragraphs[0].text
    assert "PAGE" in document.sections[0].footer._element.xml


def test_download_headers(overpass):
    xid = prepare()
    for fmt in ("pdf", "docx"):
        response = report(xid, fmt)
        assert response.headers["content-disposition"] == f'attachment; filename="geomind-research-{xid}.{fmt}"'
        assert response.headers["cache-control"] == "no-store"
    assert rr.filename('../../etc/passwd"; x', "pdf") == "geomind-research-etcpasswdx.pdf"


@pytest.mark.parametrize("body", [{"format": "html"}, {"format": "../x.pdf"}, {}, {"format": "pdf", "path": "/tmp/x"}])
def test_invalid_requests(overpass, body):
    xid = prepare()
    assert client.post(f"/api/research/executions/{xid}/report", json=body).status_code == 422


# Content correctness and evidence integrity

def test_report_contains_the_actual_evidence(overpass):
    xid = prepare()
    synthesis = synthesis_of(xid)
    text = docx_text(docx_document(report(xid, "docx").content))
    for item in synthesis.evidence:
        assert item.evidence_id in text
        assert rr.format_value(item.observed_value, item.unit) in text
    for finding in synthesis.key_findings:
        assert finding.conclusion in text
    assert synthesis.synthesis_id in text and xid in text


def test_methodology_lists_only_executed_workflows(overpass):
    xid = prepare()
    text = docx_text(docx_document(report(xid, "docx").content))
    assert "Service coverage analysis" in text
    assert "Site suitability (MCDA)" not in text and "Area comparison" not in text
    assert "local azimuthal equidistant projection" in text


def test_partial_execution_is_reported_with_missing_values(overpass):
    overpass.fail_area = "b"
    xid = prepare(COMPLEX, b=B)
    response = report(xid, "docx")
    assert response.status_code == 200
    document = docx_document(response.content)
    text = docx_text(document)
    assert "Partial: some results are incomplete" in text
    comparison = next(t for t in document.tables if t.rows[0].cells[0].text == "Metric")
    for row in comparison.rows[1:]:
        assert row.cells[2].text == "missing"  # Area B: missing, never 0
        assert row.cells[5].text == "n/a"
    areas = next(t for t in document.tables if t.rows[0].cells[0].text == "Area" and len(t.columns) == 5)
    assert areas.rows[2].cells[1].text == "not analysed"
    assert "missing, not zero" in text
    assert "No overall winner is declared" in text


def test_truncated_results_are_flagged(overpass):
    overpass.many = {"hospital"}
    xid = prepare()
    text = docx_text(docx_document(report(xid, "docx").content))
    assert "truncated" in text and "partial/truncated" in text


def test_null_values_are_preserved():
    assert rr.format_value(None, "m") == "missing (not measured)"
    assert rr.format_value(0, "count") == "0"
    assert rr._metric(None, "%") == "missing" and rr._metric(0.0, "%") == "0%"


def test_unknown_evidence_reference_is_rejected(overpass):
    xid = prepare()
    execution = ro.REGISTRY.get(xid)
    bad = synthesis_of(xid).model_copy(deep=True)
    bad.key_findings[0].evidence_ids = ["invented:evidence"]
    execution.syntheses["auto"] = bad
    response = report(xid)
    assert response.status_code == 422 and "unknown evidence" in response.json()["detail"]


def test_synthesis_from_another_execution_is_rejected(overpass):
    xid = prepare()
    execution = ro.REGISTRY.get(xid)
    execution.syntheses["auto"] = synthesis_of(xid).model_copy(update={"execution_id": "rx-other"})
    response = report(xid)
    assert response.status_code == 422 and "do not belong" in response.json()["detail"]


def test_no_fake_confidence_scores(overpass):
    xid = prepare(COMPLEX, b=B)
    text = docx_text(docx_document(report(xid, "docx").content)).lower()
    # The report may explain that no confidence is estimated, but never state one.
    assert not re.search(r"confidence[^.\n]{0,40}\d+(\.\d+)?\s*%", text)
    assert not re.search(r"\d+(\.\d+)?\s*%\s*(confidence|confident|certain)", text)


# Attribution and text safety

def test_osm_attribution_when_osm_data_was_used(overpass):
    xid = prepare()
    assert "© OpenStreetMap contributors" in pdf_text(report(xid, "pdf").content)
    assert "openstreetmap.org/copyright" in docx_text(docx_document(report(xid, "docx").content))


def test_no_osm_attribution_without_osm_data(overpass):
    xid = prepare("Create a 500 m buffer around this location")
    synthesis = synthesis_of(xid)
    assert all("OpenStreetMap" not in e.source_dataset for e in synthesis.evidence)
    text = pdf_text(report(xid, "pdf").content)
    assert "© OpenStreetMap contributors" not in text and "Buffer 500 m was created" in text
    assert overpass.calls == []


def test_user_text_is_rendered_safely(overpass):
    question = 'Identify areas with poor hospital coverage <script>alert("x")</script> & <b>bold</b> → fast\x07'
    xid = prepare(question)
    pdf = report(xid, "pdf")
    docx = report(xid, "docx")
    assert pdf.status_code == docx.status_code == 200
    text = pdf_text(pdf.content)
    assert '<script>alert("x")</script> & <b>bold</b> -> fast' in text  # literal, not markup
    assert '<script>alert("x")</script>' in docx_text(docx_document(docx.content))


def test_unsupported_glyphs_never_break_the_pdf(overpass):
    xid = prepare("Identify areas with poor hospital coverage near अस्पताल Σ")
    assert PdfReader(io.BytesIO(report(xid, "pdf").content)).pages


# Availability rules

def test_unknown_execution():
    response = report("rx-000000000000")
    assert response.status_code == 404 and "Run the research again" in response.json()["detail"]


def test_expired_execution(overpass):
    xid = prepare()
    execution = ro.REGISTRY.get(xid)
    execution.created_at -= timedelta(seconds=ro.EXECUTION_TTL_SECONDS + 1)
    response = report(xid)
    assert response.status_code == 404 and "Run the research again" in response.json()["detail"]


def test_missing_insights(overpass):
    xid = prepare(insights=False)
    response = report(xid)
    assert response.status_code == 409 and "Generate Research Insights first" in response.json()["detail"]


def test_running_execution(overpass):
    plan = make_plan(COVERAGE_Q)
    execution = ro.REGISTRY.create(plan, ro.validate_for_execution(plan))
    execution.status = "running"
    assert report(execution.id).status_code == 409


def test_failed_and_cancelled_executions(overpass):
    overpass.fail = {"hospital"}
    xid = prepare(insights=False)
    response = report(xid)
    assert response.status_code == 409 and "research failed" in response.json()["detail"]
    plan = make_plan(COVERAGE_Q)
    cancelled = ro.REGISTRY.create(plan, ro.validate_for_execution(plan))
    cancelled.close("cancelled")
    response = report(cancelled.id)
    assert response.status_code == 409 and "cancelled" in response.json()["detail"]


# No re-execution, no data fetching, no LLM

def test_report_never_reruns_or_fetches(overpass, monkeypatch):
    xid = prepare(COMPLEX, b=B)
    calls = len(overpass.calls)

    def explode(*args, **kwargs):
        raise AssertionError("must not be called while reporting")

    monkeypatch.setattr(ro, "run_execution", explode)
    monkeypatch.setattr(rs, "synthesize", explode)
    monkeypatch.setattr(research_router, "create_provider", explode)
    for fmt in ("pdf", "docx"):
        assert report(xid, fmt).status_code == 200
    assert len(overpass.calls) == calls


def test_deterministic_report_without_llm(overpass):
    xid = prepare()
    text = docx_text(docx_document(report(xid, "docx").content))
    assert "deterministic" in text and synthesis_of(xid).synthesis_mode == "deterministic"


def test_reports_are_bounded(overpass):
    overpass.many = {"hospital"}
    xid = prepare(COMPLEX, b=B)
    for fmt in ("pdf", "docx"):
        assert len(report(xid, fmt).content) < 1_000_000
    built = rr.build_report(ro.REGISTRY.get(xid), synthesis_of(xid))
    for section in built.sections:
        for block in section.blocks:
            assert len(block.rows) <= rr.MAX_TABLE_ROWS
    json.dumps([s.title for s in built.sections])
    assert all(re.match(r"^\d+\. ", s.title) for s in built.sections)
