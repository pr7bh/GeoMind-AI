"""Professional research reports (PDF and DOCX) from finished executions.

The report is assembled only from what the execution already holds: the
approved plan, the evidence digest and the M9.3 synthesis. Nothing is
re-run, nothing is fetched and no geometry is needed. Missing values stay
missing, partial results stay labelled, and no confidence scores are added.

A format-neutral `Report` is built once, then rendered by ReportLab (PDF)
or python-docx (DOCX), entirely in memory.
"""

import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any
from xml.sax.saxutils import escape

from app.schemas.research import EvidenceItem, ResearchSynthesis
from app.services.workflows import catalog

REPORT_VERSION = "GeoMind AI research report 1.0"
OSM_ATTRIBUTION = (
    "© OpenStreetMap contributors. Data available under the Open Database License (ODbL): "
    "https://www.openstreetmap.org/copyright"
)
MAX_TABLE_ROWS = 200
GENERAL_LIMITATIONS_TITLE = "General methodological limitations (apply to every analysis of this kind)"

# Methodology actually used by each capability (only executed ones are shown).
CAPABILITY_METHODS: dict[str, list[str]] = {
    "coverage": catalog.COVERAGE_METHODOLOGY,
    "comparison": catalog.COMPARISON_METHODOLOGY,
    "suitability": catalog.SUITABILITY_METHODOLOGY,
    "nearby_count": [
        "Features of the source category are discovered around Area A; features of the target category are then "
        "searched within the stated radius of each source.",
        "Each source is ranked by the number of distinct target features within the radius; equal counts share a rank.",
        "Sources whose nearby search failed are listed separately and never counted as zero.",
    ],
    "discover_features": [
        "Mapped OpenStreetMap features of one category are retrieved within the stated radius of Area A; distances "
        "are geodesic (WGS84) from Area A to each feature.",
    ],
    "find_nearest_features": [
        "Discovered features are ranked by geodesic (WGS84) distance from Area A.",
    ],
    "create_buffer": ["A circular buffer of the stated distance is drawn around Area A."],
    "spatial_query": [
        "Discovered features are kept when they lie within (or intersect) a circular buffer around Area A.",
    ],
}
PARAMETER_LABELS = {
    "category": "Service category",
    "facility": "Facility type",
    "radius_km": "Study radius (km)",
    "search_radius_km": "Discovery search radius (km)",
    "cell_size_m": "Grid cell size (m)",
    "well_covered_m": "Well-covered threshold (m)",
    "moderate_m": "Moderate threshold (m)",
    "limit": "Maximum features per search",
    "normalization": "Score normalisation",
    "max_candidates": "Maximum candidates",
    "pharmacy_direction": "Pharmacy proximity direction",
    "existing_saturation_m": "Distance saturation (m)",
    "distance_metres": "Buffer distance (m)",
    "operation": "Spatial relationship",
    "weights": "Criterion weights",
    "area": "Area",
}


class ReportError(Exception):
    """The report cannot be produced; the message is safe to show."""

    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


# Format-neutral report structure

@dataclass
class Block:
    kind: str  # paragraph | muted | warning | bullets | table | subheading | keyvalue
    text: str = ""
    items: list[str] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    # Cells may contain "\n": the first line is the main text, later lines are details.
    rows: list[list[str]] = field(default_factory=list)
    widths: list[float] = field(default_factory=list)
    # First column holds identifiers (rendered monospaced).
    id_column: bool = False


@dataclass
class Section:
    title: str
    blocks: list[Block]


@dataclass
class Report:
    title: str
    question: str
    cover: list[tuple[str, str]]
    sections: list[Section]
    execution_id: str
    generated_at: datetime
    status_note: str
    # OpenStreetMap attribution is shown when OSM data contributed.
    uses_osm: bool


def _p(text: str) -> Block:
    return Block("paragraph", text=text)


def _muted(text: str) -> Block:
    return Block("muted", text=text)


def _warn(text: str) -> Block:
    return Block("warning", text=text)


def _bullets(items: list[str]) -> Block:
    return Block("bullets", items=[i for i in items if i])


def _sub(text: str) -> Block:
    return Block("subheading", text=text)


def _table(headers: list[str], rows: list[list[str]], widths: list[float], id_column: bool = False) -> Block:
    return Block("table", headers=headers, rows=rows[:MAX_TABLE_ROWS], widths=widths, id_column=id_column)


# Values: missing stays missing, units are kept

def format_value(value: Any, unit: str | None) -> str:
    if value is None:
        return "missing (not measured)"
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, str):
        return value
    if unit == "m":
        return f"{value:g} m ({value / 1000:.2f} km)"
    if unit == "%":
        return f"{value:g}%"
    if unit == "weight":
        return f"{value:g} ({value * 100:.0f}%)"
    if unit in (None, "count", "class"):
        return f"{value:g}"
    return f"{value:g} {unit}"


def _metric(value: Any, unit: str) -> str:
    if value is None:
        return "missing"
    if unit == "m":
        return f"{value / 1000:.2f} km"
    if unit == "%":
        return f"{value:g}%"
    if unit == "per km²":
        return f"{value:.3f}/km²"
    return f"{value:g}"


def _coords(coords: list[float] | None) -> str:
    return f"{coords[1]:.5f}, {coords[0]:.5f} (lat, lon, WGS84)" if coords else "not selected"


# Validation

def validate_source(execution_id: str, plan_id: str, synthesis: ResearchSynthesis) -> None:
    """The synthesis must belong to this execution and every reference must resolve."""
    if synthesis.execution_id != execution_id or synthesis.plan_id != plan_id:
        raise ReportError("The research insights do not belong to this execution.")
    ids = [e.evidence_id for e in synthesis.evidence]
    if len(ids) != len(set(ids)):
        raise ReportError("The research insights contain duplicate evidence ids.")
    known = set(ids)
    if any(e.source_execution_id != execution_id for e in synthesis.evidence):
        raise ReportError("Some evidence does not come from this execution.")
    referenced = (
        [i for f in synthesis.key_findings for i in f.evidence_ids]
        + [i for c in synthesis.comparisons for i in c.evidence_ids]
        + [i for c in synthesis.candidate_insights for i in c.evidence_ids]
        + [i for r in synthesis.recommendations for i in r.evidence_ids]
    )
    missing = sorted(set(referenced) - known)
    if missing:
        raise ReportError(f"Findings cite unknown evidence ({', '.join(missing[:5])}).")
    if any(not f.evidence_ids for f in synthesis.key_findings):
        raise ReportError("A finding has no supporting evidence.")


# Report assembly

def build_report(execution, synthesis: ResearchSynthesis, now: datetime | None = None) -> Report:
    plan, digest = execution.plan, execution.digest
    if digest is None:
        raise ReportError("No evidence was kept for this execution. Run the research again.", 409)
    validate_source(execution.id, plan.id, synthesis)
    now = now or datetime.now(timezone.utc)
    evidence = {e.evidence_id: e for e in synthesis.evidence}
    capabilities = digest["capabilities"]
    # Every capability that ran (including ones that failed) used its methodology.
    executed = list(capabilities)

    status_note = {
        "complete": "Complete: all reported results come from completed steps with complete data.",
        "partial": "Partial: some results are incomplete. Read the data-quality section before relying on them.",
    }.get(synthesis.overall_status, synthesis.overall_status)

    sections: list[Section] = []

    # 1. Executive summary
    blocks = [_sub("Research objective"), _p(plan.objective), _sub("Summary"), _p(synthesis.executive_summary)]
    if synthesis.key_findings:
        blocks += [_sub("Most important findings"), _bullets([
            f"{f.conclusion} [evidence: {', '.join(f.evidence_ids[:4])}{'…' if len(f.evidence_ids) > 4 else ''}]"
            for f in synthesis.key_findings[:4]
        ])]
    if synthesis.overall_status != "complete" or synthesis.data_quality_warnings:
        blocks.append(_warn(status_note))
        blocks += [_warn(w) for w in synthesis.data_quality_warnings[:5]]
    sections.append(Section("Executive Summary", blocks))

    # 2. Research objectives and parameters
    areas = [a for a in plan.selected_areas if a.coordinates or a.required]
    blocks = [
        _sub("Research question"), _p(plan.original_question),
        _sub("Areas analysed"),
        _table(["Area", "Location"], [[a.label, _coords(a.coordinates)] for a in areas], [0.35, 0.65]),
        _sub("Capabilities executed"),
        _bullets([f"{catalog.CAPABILITIES[c].title}: {catalog.CAPABILITIES[c].summary}" for c in executed if c in catalog.CAPABILITIES]),
        _sub("Parameters (as executed)"),
        _table(["Step", "Parameter", "Value"], _parameter_rows(plan), [0.32, 0.33, 0.35]),
    ]
    if plan.interpretation:
        blocks += [_sub("Scope"), _bullets(plan.interpretation)]
    if plan.assumptions:
        blocks += [_sub("Assumptions"), _bullets(plan.assumptions)]
    sections.append(Section("Research Objectives", blocks))

    # 3. Methodology (executed capabilities only)
    blocks = []
    for capability in executed:
        lines = CAPABILITY_METHODS.get(capability)
        if lines:
            blocks += [_sub(catalog.CAPABILITIES[capability].title), _bullets(lines)]
    blocks += [
        _sub("Coordinates and units"),
        _bullets([
            "Inputs and outputs use WGS84 longitude/latitude. Grid and distance calculations for coverage and "
            "suitability use a local azimuthal equidistant projection centred on the study location.",
            "Distances are in metres (shown in kilometres where helpful); percentages refer to shares of grid cells.",
        ]),
        _sub("Synthesis rules"), _bullets(synthesis.methodology),
    ]
    sections.append(Section("Methodology", blocks))

    # 4. Key findings
    blocks = []
    if not synthesis.key_findings:
        blocks.append(_muted("No evidence-supported findings."))
    for finding in synthesis.key_findings:
        source = "AI-written, evidence-checked" if finding.source == "llm" else "Deterministic"
        blocks += [
            _sub(f"{finding.finding_id} · {finding.title}"),
            _p(finding.conclusion),
            _muted(f"Quality: {finding.quality} · Scope: {finding.geographic_scope} · Source: {source}"),
        ]
        if finding.explanation:
            blocks.append(_p(finding.explanation))
        blocks += [_warn(c) for c in finding.caveats]
        blocks.append(_table(
            ["Evidence", "Measurement", "Value", "Type"],
            [[e, evidence[e].title, format_value(evidence[e].observed_value, evidence[e].unit), evidence[e].kind]
             for e in finding.evidence_ids],
            [0.3, 0.34, 0.22, 0.14],
            id_column=True,
        ))
    sections.append(Section("Key Findings", blocks))

    # 5. Comparative analysis
    if synthesis.comparisons:
        sections.append(Section("Comparative Analysis", _comparison_blocks(synthesis, digest)))

    # 6. Site suitability
    if synthesis.candidate_insights:
        sections.append(Section("Site Suitability", _suitability_blocks(synthesis, digest)))

    # Nearby-count ranking (hospital–pharmacy style)
    nearby = capabilities.get("nearby_count")
    if nearby and nearby.get("ranking"):
        sections.append(Section("Nearby-Count Ranking", _nearby_blocks(nearby)))

    # 7. Evidence register
    sections.append(Section("Evidence Register", [
        _muted("Every value below is read from this execution's workflow results (measured) or calculated from them "
               "(derived). Missing values are shown as missing, never as zero."),
        _table(
            ["Evidence ID", "Finding or measurement", "Value", "Type", "Source", "Quality"],
            [_evidence_row(e) for e in synthesis.evidence],
            [0.2, 0.3, 0.13, 0.12, 0.14, 0.11],
            id_column=True,
        ),
    ]))

    # 8. Data quality and limitations
    blocks = [_sub("Execution-specific warnings")]
    specific = list(dict.fromkeys(synthesis.data_quality_warnings + digest["execution"].get("errors", [])))
    blocks += [_warn(w) for w in specific] if specific else [_muted("No data-quality warnings were raised for this execution.")]
    truncated = [e.evidence_id for e in synthesis.evidence if e.partial_or_truncated]
    if truncated:
        blocks.append(_warn(f"{len(truncated)} evidence item(s) are partial or truncated: {', '.join(truncated[:8])}"
                            f"{'…' if len(truncated) > 8 else ''}. Counts and rankings based on them may be incomplete."))
    if synthesis.unsupported_questions:
        blocks += [_sub("Not analysed"), _bullets(synthesis.unsupported_questions)]
    blocks += [_sub(GENERAL_LIMITATIONS_TITLE), _bullets(synthesis.limitations)]
    blocks += [_sub("Data freshness"), _bullets(_freshness(digest, execution))]
    sections.append(Section("Data Quality and Limitations", blocks))

    # 9. Conclusions and next steps
    blocks = [_sub("Evidence-supported conclusions")]
    blocks.append(_bullets([f"{f.conclusion} [{', '.join(f.evidence_ids[:3])}]" for f in synthesis.key_findings])
                  if synthesis.key_findings else _muted("No conclusions are supported by the available evidence."))
    unresolved = list(synthesis.unsupported_questions)
    unresolved += [f"{c.metric}: {c.note}" for c in synthesis.comparisons if not c.comparable]
    if unresolved:
        blocks += [_sub("Unresolved questions"), _bullets(unresolved)]
    if synthesis.recommendations:
        blocks += [_sub("Recommended next steps"), _bullets([
            f"{r.text} [{', '.join(r.evidence_ids)}]" + (" (AI-written, evidence-checked)" if r.source == "llm" else "")
            for r in synthesis.recommendations
        ])]
    sections.append(Section("Conclusions and Next Steps", blocks))

    # 10. Sources and attribution
    datasets = sorted({e.source_dataset for e in synthesis.evidence})
    blocks = []
    if any("OpenStreetMap" in d for d in datasets):
        blocks += [_sub("OpenStreetMap"), _p(OSM_ATTRIBUTION),
                   _muted("Retrieved through the Overpass API at the time of execution.")]
    if datasets:
        blocks += [_sub("Datasets cited by the evidence"), _bullets(datasets)]
    blocks += [
        _sub("Report metadata"),
        _table(["Field", "Value"], [
            ["Report generated", now.strftime("%Y-%m-%d %H:%M:%S UTC")],
            ["Report version", REPORT_VERSION],
            ["Execution ID", execution.id],
            ["Execution status", execution.status],
            ["Execution finished", execution.finished_at.strftime("%Y-%m-%d %H:%M:%S UTC") if execution.finished_at else "n/a"],
            ["Plan ID", plan.id],
            ["Insights ID", synthesis.synthesis_id],
            ["Insights mode", synthesis.synthesis_mode],
            ["Insights generated", synthesis.generated_at.strftime("%Y-%m-%d %H:%M:%S UTC")],
        ], [0.35, 0.65]),
        _muted("Map imagery is not included: reports are built from retained evidence, which deliberately excludes "
               "geometry. Add the result layers to the GeoMind map to view them."),
    ]
    sections.append(Section("Sources and Attribution", blocks))

    for number, section in enumerate(sections, start=1):
        section.title = f"{number}. {section.title}"

    return Report(
        title=synthesis.title,
        question=plan.original_question,
        cover=[
            ("Research question", plan.original_question),
            ("Generated", now.strftime("%Y-%m-%d %H:%M UTC")),
            ("Execution ID", execution.id),
            ("Execution status", execution.status.replace("_", " ")),
            ("Evidence status", synthesis.overall_status.replace("_", " ")),
            ("Insights mode", synthesis.synthesis_mode.replace("_", " ")),
            ("Report version", REPORT_VERSION),
        ],
        sections=sections,
        execution_id=execution.id,
        generated_at=now,
        status_note=status_note,
        uses_osm=any("OpenStreetMap" in d for d in datasets),
    )


def _parameter_rows(plan) -> list[list[str]]:
    rows = []
    for step in plan.steps:
        for key, value in step.arguments.items():
            if key not in PARAMETER_LABELS or value is None:
                continue
            if isinstance(value, dict):
                value = ", ".join(f"{k.replace('_', ' ')} {v:g}" for k, v in value.items())
            elif isinstance(value, float):
                value = f"{value:g}"
            rows.append([f"{step.execution_order}. {step.title}", PARAMETER_LABELS[key], str(value)])
    return rows


def _comparison_blocks(synthesis: ResearchSynthesis, digest: dict[str, Any]) -> list[Block]:
    comparison = digest["capabilities"].get("comparison") or {}
    rows_by_label = {r["label"]: r for r in comparison.get("metrics", [])}
    blocks = []
    areas = comparison.get("areas") or {}
    area_rows = []
    for key in ("a", "b"):
        area = areas.get(key) or {}
        metrics = area.get("metrics")
        if metrics:
            area_rows.append([f"Area {key.upper()}", _coords(metrics["centre"]), f"{metrics['radius_km']:g} km",
                              f"{metrics['cell_size_m']} m", str(metrics["total_cells"])])
        else:
            area_rows.append([f"Area {key.upper()}", "not analysed", "missing", "missing",
                              f"missing ({area.get('error') or 'did not run'})"])
    blocks += [_sub("Study areas"), _table(["Area", "Centre", "Radius", "Cell size", "Cells"], area_rows,
                                            [0.12, 0.38, 0.14, 0.14, 0.22])]
    rows = []
    for c in synthesis.comparisons:
        relative = rows_by_label.get(c.metric, {}).get("relative_difference_pct")
        rows.append([
            c.metric, _metric(c.a_value, c.unit), _metric(c.b_value, c.unit),
            "missing" if c.difference is None else _metric(c.difference, c.unit) if c.unit != "%" else f"{c.difference:g} pp",
            "n/a" if relative is None or c.unit == "%" else f"{relative:g}%",
            {"a": "A", "b": "B", "tie": "Equal"}.get(c.better or "", "n/a"),
            c.note,
        ])
    blocks += [
        _sub("Metric by metric"),
        _table(["Metric", "Area A", "Area B", "A − B", "Relative", "Better", "Note"], rows,
               [0.2, 0.11, 0.11, 0.11, 0.1, 0.08, 0.29]),
        _p("No overall winner is declared: each metric is compared on its own, and missing values are not treated as zero."),
    ]
    return blocks


def _suitability_blocks(synthesis: ResearchSynthesis, digest: dict[str, Any]) -> list[Block]:
    suit = digest["capabilities"].get("suitability") or {}
    blocks = [_warn("Candidates are analytical grid points, not verified or available land parcels. Land availability, "
                    "zoning, ownership and access were not analysed.")]
    if suit:
        blocks += [_sub("Configuration"), _table(["Parameter", "Value"], [
            ["Facility", str(suit.get("facility_label"))],
            ["Study area", f"{suit.get('radius_km'):g} km around {_coords(suit.get('centre'))}"],
            ["Candidates", f"{suit.get('scored_count')} scored, {suit.get('unscored_count')} unscored of {suit.get('candidate_count')}"],
            ["Normalisation", str(suit.get("normalization"))],
        ], [0.3, 0.7])]
        blocks += [_sub("Criteria and weights"), _table(
            ["Criterion", "Status", "Weight", "Direction", "Note"],
            [[c["name"], c["status"], f"{c['weight'] * 100:.0f}%",
              "higher is better" if c["higher_is_better"] else "lower is better", c.get("note") or ""]
             for c in suit.get("criteria", [])],
            [0.32, 0.13, 0.1, 0.17, 0.28])]
        if (suit.get("tied_at_top") or 1) > 1:
            blocks.append(_warn(f"{suit['tied_at_top']} candidates share the highest score; the ranking does not separate them."))
    blocks += [_sub("Top candidates"), _table(
        ["Rank", "Candidate", "Location", "Score", "Coverage cell"],
        [[f"{c.rank}{'=' if c.tied else ''}", c.candidate_id, _coords(c.coordinates), f"{c.score:g}",
          (c.coverage_class or "not available").replace("_", " ")] for c in synthesis.candidate_insights],
        [0.09, 0.15, 0.4, 0.12, 0.24])]
    top = {c["id"]: c for c in suit.get("top", [])}
    criteria = {c["id"]: c for c in suit.get("criteria", [])}
    for insight in synthesis.candidate_insights:
        blocks += [_sub(f"Candidate {insight.candidate_id}: score breakdown"), _p(insight.explanation)]
        detail = top.get(insight.candidate_id)
        if detail:
            blocks.append(_table(
                ["Criterion", "Measured value", "Score (0–100)", "Weight", "Points"],
                [[criteria.get(cid, {}).get("name", cid), format_value(part["raw"], criteria.get(cid, {}).get("unit"))
                  + (" (lower bound)" if part.get("lower_bound") else ""),
                  "missing" if part["normalized"] is None else f"{part['normalized']:g}",
                  f"{part['weight'] * 100:.0f}%", "missing" if part["contribution"] is None else f"{part['contribution']:g}"]
                 for cid, part in detail["criteria"].items()],
                [0.32, 0.24, 0.16, 0.12, 0.16]))
        if insight.strengths:
            blocks.append(_bullets([f"Strength: {s}" for s in insight.strengths]))
        if insight.weaknesses:
            blocks.append(_bullets([f"Weakness: {w}" for w in insight.weaknesses]))
    return blocks


def _nearby_blocks(nearby: dict[str, Any]) -> list[Block]:
    target = ((nearby.get("nearby") or {}).get("label") or "features").lower()
    rows = [[f"{r['rank']}{'=' if r['tied'] else ''}", r["name"],
             f"{r['nearby_count']}{'+' if r['count_is_lower_bound'] else ''}",
             "n/a" if r.get("nearest_m") is None else f"{r['nearest_m'] / 1000:.2f} km"]
            for r in nearby["ranking"]]
    blocks = [_table(["Rank", "Feature", f"{target.capitalize()} nearby", "Nearest"], rows, [0.1, 0.5, 0.2, 0.2])]
    if nearby.get("unranked"):
        blocks += [_warn("Not ranked (search failed; counts unknown, not zero):"),
                   _bullets([f"{u['name']}: {u['reason']}" for u in nearby["unranked"]])]
    return blocks


def _evidence_row(item: EvidenceItem) -> list[str]:
    details = [f"Scope: {item.geographic_scope}", f"Method: {item.methodology}"]
    details += [f"Limitation: {line}" for line in item.limitations]
    quality = item.quality + (" · partial/truncated" if item.partial_or_truncated else "")
    return [
        item.evidence_id,
        item.title + "\n" + "\n".join(details),
        format_value(item.observed_value, item.unit),
        f"{item.kind}\n{item.finding_type.replace('_', ' ')}",
        f"{item.capability_id}\n{item.source_step_id or '—'}",
        quality,
    ]


def _freshness(digest: dict[str, Any], execution) -> list[str]:
    lines = []
    stamps = set()
    for entry in digest["capabilities"].values():
        services = entry.get("services")
        if isinstance(services, dict):
            for value in [services] + [v for v in services.values() if isinstance(v, dict)]:
                if value.get("osm_base"):
                    stamps.add(value["osm_base"])
    if stamps:
        lines.append("OpenStreetMap data as of: " + ", ".join(sorted(stamps)) + ".")
    lines.append("This report is a snapshot of one execution. GeoMind keeps executions in memory for 30 minutes and "
                 "loses them when the server restarts; re-running the research may give different results if "
                 "OpenStreetMap data has changed.")
    return lines


# Text safety

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text: Any) -> str:
    """Remove characters that are invalid in XML documents."""
    return _CONTROL.sub("", str(text))


def filename(execution_id: str, fmt: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "", execution_id)[:40] or "execution"
    return f"geomind-research-{safe}.{fmt}"


# PDF (ReportLab)

_PDF_FONTS: dict[str, Any] = {}
_PDF_REPLACEMENTS = {"→": "->", "←": "<-", "Σ": "sum", "“": '"', "”": '"', "’": "'", "‘": "'"}


def _pdf_fonts() -> dict[str, Any]:
    """Bitstream Vera (bundled with ReportLab): broader Unicode than Helvetica."""
    if _PDF_FONTS:
        return _PDF_FONTS
    import reportlab
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    folder = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    for name, file in (("GeoVera", "Vera.ttf"), ("GeoVera-Bold", "VeraBd.ttf"), ("GeoVera-Italic", "VeraIt.ttf"),
                       ("GeoVera-BoldItalic", "VeraBI.ttf")):
        pdfmetrics.registerFont(TTFont(name, os.path.join(folder, file)))
    addMapping("GeoVera", 0, 0, "GeoVera")
    addMapping("GeoVera", 1, 0, "GeoVera-Bold")
    addMapping("GeoVera", 0, 1, "GeoVera-Italic")
    addMapping("GeoVera", 1, 1, "GeoVera-BoldItalic")
    _PDF_FONTS["glyphs"] = set(pdfmetrics.getFont("GeoVera").face.charToGlyph)
    return _PDF_FONTS


def _pdf_text(text: Any) -> str:
    """Escape for ReportLab's markup and replace glyphs the font lacks."""
    glyphs = _pdf_fonts()["glyphs"]
    out = []
    for char in clean_text(text):
        char = _PDF_REPLACEMENTS.get(char, char)
        if len(char) == 1 and ord(char) not in glyphs and char not in "\n\t":
            ascii_form = unicodedata.normalize("NFKD", char).encode("ascii", "ignore").decode()
            char = ascii_form or "?"
        out.append(char)
    return escape("".join(out))


def render_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import (
        KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle,
    )

    _pdf_fonts()
    emerald, ink, muted, line = colors.HexColor("#059669"), colors.HexColor("#111827"), colors.HexColor("#6b7280"), colors.HexColor("#e5e7eb")
    amber, amber_bg, header_bg = colors.HexColor("#92400e"), colors.HexColor("#fffbeb"), colors.HexColor("#ecfdf5")

    base = ParagraphStyle("base", fontName="GeoVera", fontSize=9.5, leading=13.5, textColor=ink, alignment=TA_LEFT)
    styles = {
        "title": ParagraphStyle("title", parent=base, fontName="GeoVera-Bold", fontSize=22, leading=27, spaceAfter=10),
        "brand": ParagraphStyle("brand", parent=base, fontName="GeoVera-Bold", fontSize=11, textColor=emerald, spaceAfter=24),
        "h1": ParagraphStyle("h1", parent=base, fontName="GeoVera-Bold", fontSize=14.5, leading=19, textColor=ink,
                             spaceBefore=4, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base, fontName="GeoVera-Bold", fontSize=10.5, leading=14, textColor=emerald,
                             spaceBefore=8, spaceAfter=3),
        "paragraph": ParagraphStyle("p", parent=base, spaceAfter=5),
        "muted": ParagraphStyle("muted", parent=base, fontSize=8.5, leading=12, textColor=muted, spaceAfter=5),
        "warning": ParagraphStyle("warning", parent=base, fontSize=8.8, leading=12.5, textColor=amber, backColor=amber_bg,
                                  borderPadding=(4, 6, 4, 6), spaceBefore=3, spaceAfter=7),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=12, bulletIndent=2, spaceAfter=2.5),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8, leading=10.5),
        "cellhead": ParagraphStyle("cellhead", parent=base, fontName="GeoVera-Bold", fontSize=8, leading=10.5, textColor=emerald),
        "cellid": ParagraphStyle("cellid", parent=base, fontName="Courier", fontSize=6.8, leading=9.5),
    }

    def cell(text: str, head: bool = False, identifier: bool = False):
        lines = str(text).split("\n")
        markup = _pdf_text(lines[0])
        if len(lines) > 1:
            markup += "".join(f'<br/><font size="7" color="#6b7280">{_pdf_text(l)}</font>' for l in lines[1:])
        return Paragraph(markup, styles["cellhead" if head else "cellid" if identifier else "cell"])

    page_width = A4[0] - 36 * mm
    story: list[Any] = [
        Spacer(1, 30 * mm),
        Paragraph("GeoMind AI · Research Report", styles["brand"]),
        Paragraph(_pdf_text(report.title), styles["title"]),
        Spacer(1, 4 * mm),
    ]
    cover = LongTable(
        [[cell(k, True), cell(v)] for k, v in report.cover],
        colWidths=[page_width * 0.28, page_width * 0.72],
    )
    cover.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [cover, Spacer(1, 8 * mm), Paragraph(_pdf_text(report.status_note), styles["muted"]), PageBreak()]

    for index, section in enumerate(report.sections):
        if index:
            story.append(Spacer(1, 6 * mm))
        heading = Paragraph(_pdf_text(section.title), styles["h1"])
        first = True
        for block in section.blocks:
            flowables = _pdf_block(block, styles, cell, page_width, LongTable, TableStyle, header_bg, line, Paragraph)
            if first:
                # Keep a heading with its first block so it never ends a page.
                story.append(KeepTogether([heading, flowables[0]]))
                story += flowables[1:]
                first = False
            else:
                story += flowables
        if first:
            story.append(heading)

    title = report.title

    class NumberedCanvas(pdfcanvas.Canvas):
        """Header, footer and "Page N of M" (two-pass page count)."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                if self._pageNumber > 1:
                    self.setFont("GeoVera", 7.5)
                    self.setFillColor(muted)
                    self.drawString(18 * mm, A4[1] - 12 * mm, _pdf_text(f"GeoMind AI · {title}")[:110])
                    self.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {self._pageNumber} of {total}")
                    self.drawString(18 * mm, 10 * mm, f"Execution {report.execution_id}"
                                    + (" · © OpenStreetMap contributors (ODbL)" if report.uses_osm else ""))
                    self.setStrokeColor(line)
                    self.line(18 * mm, A4[1] - 14 * mm, A4[0] - 18 * mm, A4[1] - 14 * mm)
                super().showPage()
            super().save()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm, bottomMargin=18 * mm,
        title=clean_text(report.title)[:200], author="GeoMind AI", subject="GIS research report",
        creator=REPORT_VERSION,
    )
    doc.build(story, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


def _pdf_block(block, styles, cell, page_width, LongTable, TableStyle, header_bg, line, Paragraph) -> list[Any]:
    if block.kind == "table":
        if not block.rows:
            return [Paragraph(_pdf_text("No data."), styles["muted"])]
        table = LongTable(
            [[cell(h, True) for h in block.headers]]
            + [[cell(c, identifier=block.id_column and i == 0) for i, c in enumerate(row)] for row in block.rows],
            colWidths=[page_width * w for w in block.widths], repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), header_bg),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        return [table, Paragraph("", styles["muted"])]
    if block.kind == "bullets":
        return [Paragraph(_pdf_text(item), styles["bullet"], bulletText="•") for item in block.items] or \
            [Paragraph("", styles["muted"])]
    style = {"subheading": "h2", "muted": "muted", "warning": "warning"}.get(block.kind, "paragraph")
    return [Paragraph(_pdf_text(block.text), styles[style])]


# DOCX (python-docx)

def render_docx(report: Report) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    emerald, muted, amber = RGBColor(0x05, 0x96, 0x69), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x92, 0x40, 0x0E)
    document = Document()
    document.core_properties.title = clean_text(report.title)[:200]
    document.core_properties.author = "GeoMind AI"
    document.core_properties.subject = "GIS research report"
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    for level, size in ((1, 15), (2, 11.5)):
        style = document.styles[f"Heading {level}"]
        style.font.color.rgb = emerald if level == 2 else RGBColor(0x11, 0x18, 0x27)
        style.font.size = Pt(size)

    def run(paragraph, text, *, bold=False, color=None, size=None, italic=False):
        r = paragraph.add_run(clean_text(text))
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
        return r

    def shade(cell_, fill):
        properties = cell_._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    def table(headers, rows):
        t = document.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        header = t.rows[0]
        # Repeat the header row on every page.
        header_properties = header._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_properties.append(repeat)
        for c, text in zip(header.cells, headers):
            c.paragraphs[0].text = ""
            run(c.paragraphs[0], text, bold=True, color=emerald, size=8.5)
            shade(c, "ECFDF5")
        for values in rows:
            cells = t.add_row().cells
            for c, value in zip(cells, values):
                lines = str(value).split("\n")
                run(c.paragraphs[0], lines[0], size=8.5)
                for extra in lines[1:]:
                    run(c.add_paragraph(), extra, size=7.5, color=muted)
        return t

    # Cover
    brand = document.add_paragraph()
    run(brand, "GeoMind AI · Research Report", bold=True, color=emerald, size=11)
    heading = document.add_paragraph()
    run(heading, report.title, bold=True, size=22)
    table(["Field", "Value"], [[k, v] for k, v in report.cover])
    note = document.add_paragraph()
    run(note, report.status_note, color=muted, size=9, italic=True)
    note.add_run().add_break(WD_BREAK.PAGE)

    section = document.sections[0]
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    run(header, f"GeoMind AI · {report.title}", color=muted, size=8)
    footer = section.footer.paragraphs[0]
    attribution = " · © OpenStreetMap contributors (ODbL)" if report.uses_osm else ""
    run(footer, f"Execution {report.execution_id}{attribution} · Page ", color=muted, size=8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = footer.add_run()
        r.font.size, r.font.color.rgb = Pt(8), muted
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        r._r.append(element)

    for sec in report.sections:
        document.add_heading(clean_text(sec.title), level=1)
        for block in sec.blocks:
            if block.kind == "subheading":
                document.add_heading(clean_text(block.text), level=2)
            elif block.kind == "table":
                if block.rows:
                    table(block.headers, block.rows)
                    document.add_paragraph()
                else:
                    run(document.add_paragraph(), "No data.", color=muted, size=9)
            elif block.kind == "bullets":
                for item in block.items:
                    run(document.add_paragraph(style="List Bullet"), item)
            elif block.kind == "warning":
                run(document.add_paragraph(), "⚠ " + block.text, color=amber, size=9.5)
            elif block.kind == "muted":
                run(document.add_paragraph(), block.text, color=muted, size=9, italic=True)
            else:
                run(document.add_paragraph(), block.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def render(report: Report, fmt: str) -> bytes:
    return render_pdf(report) if fmt == "pdf" else render_docx(report)
